"""Build backlog/user_stories.csv (Jira import) and docs/P2P_Requirements_Workbook.xlsx."""
from pathlib import Path
import csv
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
C = json.loads((ROOT / "scripts" / "content.json").read_text(encoding="utf-8"))

# ------------------------------------------------------------------ Jira CSV
(ROOT / "backlog").mkdir(exist_ok=True)
with open(ROOT / "backlog" / "user_stories.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh)
    w.writerow(["Issue Id", "Issue Type", "Summary", "Description", "Acceptance Criteria", "Priority", "Story Points", "Epic Name", "Labels", "Requirements"])
    for s in C["stories"]:
        reqs = ", ".join(r["id"] for r in C["requirements"] if s["id"] in r["story"])
        w.writerow([s["id"], "Story", f"{s['want'][0].upper()}{s['want'][1:]}",
                    f"As a {s['role']}, I want {s['want']}, so that {s['so']}.",
                    "\n".join(f"- {a}" for a in s["ac"]),
                    {"Must": "Highest", "Should": "High", "Could": "Medium"}[s["priority"]],
                    s["points"], s["epic"], f"MoSCoW-{s['priority']}", reqs])

# readable Markdown copy of the backlog for GitHub
md = ["# User story backlog", "", "Import `user_stories.csv` into Jira (or Azure DevOps) to recreate this backlog.", ""]
for epic in dict.fromkeys(s["epic"] for s in C["stories"]):
    md += [f"## {epic}", ""]
    for s in [x for x in C["stories"] if x["epic"] == epic]:
        reqs = ", ".join(r["id"] for r in C["requirements"] if s["id"] in r["story"])
        md += [f"### {s['id']} · {s['priority']} · {s['points']} points", "",
               f"As a **{s['role']}**, I want {s['want']}, so that {s['so']}.", "", "**Acceptance criteria**", ""]
        md += [f"- {a}" for a in s["ac"]]
        md += ["", f"Traces to: {reqs}", ""]
(ROOT / "backlog" / "user_stories.md").write_text("\n".join(md), encoding="utf-8")

# ------------------------------------------------------------------ workbook
ARIAL = "Arial"
HEAD = PatternFill("solid", fgColor="1F3A68")
HFONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
BODY = Font(name=ARIAL, size=10)
INPUT = Font(name=ARIAL, size=10, color="0000FF")
THIN = Side(style="thin", color="C4CCD6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def sheet(title, headers, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEAD, HFONT, Alignment(wrap_text=True, vertical="center"), BOX
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BODY, WRAP, BOX
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


# Read me
ws = wb.active
ws.title = "Read me"
lines = [
    ("Procure-to-pay requirements workbook", Font(name=ARIAL, bold=True, size=14)),
    ("Portfolio case study by Chiamaka Ikpo. Organisation anonymised; process generalised; values illustrative.", BODY),
    ("", BODY),
    ("Tabs", Font(name=ARIAL, bold=True, size=11)),
    ("Summary – requirement and test coverage, calculated with formulas from the other tabs", BODY),
    ("Stakeholders – stakeholder register with interest, influence and engagement approach", BODY),
    ("RACI – who is Responsible, Accountable, Consulted and Informed for each activity", BODY),
    ("Requirements – traceability: pain point → requirement → user story → UAT test cases", BODY),
    ("UAT Tests – test cases for user acceptance testing; update the Status column during testing", BODY),
    ("", BODY),
    ("How to use", Font(name=ARIAL, bold=True, size=11)),
    ("Blue text = cells you edit (test status, tester, date). Black text = reference content or formulas.", BODY),
    ("Status values: Not run, Pass, Fail, Blocked. Summary updates automatically.", BODY),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions["A"].width = 110

# Stakeholders
sheet("Stakeholders", ["ID", "Stakeholder", "Interest", "Influence", "Impact on them", "Engagement approach", "RACI role"],
      C["stakeholders"], [8, 26, 40, 11, 14, 40, 14])

# RACI
ws = sheet("RACI", ["Activity", *C["raci"]["roles"]], C["raci"]["rows"], [30, 14, 14, 14, 14])
for row in ws.iter_rows(min_row=2, min_col=2):
    for c in row:
        c.alignment = Alignment(horizontal="center", vertical="top")
ws.conditional_formatting.add("B2:E7", FormulaRule(formula=['ISNUMBER(SEARCH("A",B2))'], fill=PatternFill("solid", fgColor="DCE6F7")))
r = len(C["raci"]["rows"]) + 3
ws.cell(row=r, column=1, value="Check: exactly one Accountable per activity").font = Font(name=ARIAL, bold=True, size=10)
for i in range(2, 2 + len(C["raci"]["rows"])):
    ws.cell(row=i, column=7, value=f'=IF(COUNTIF(B{i}:E{i},"*A*")=1,"OK","Fix: "&COUNTIF(B{i}:E{i},"*A*")&" A")').font = BODY
ws.cell(row=1, column=7, value="One 'A'?").font = HFONT
ws.cell(row=1, column=7).fill = HEAD
ws.column_dimensions["G"].width = 12

# UAT tests: two per story (one per acceptance criterion, max 2)
tests = []
n = 1
for s in C["stories"]:
    for ac in s["ac"][:2]:
        tests.append([f"TC-{n:02d}", s["id"], ac, "Not run", "", ""])
        n += 1
sheet("UAT Tests", ["Test ID", "Story", "Acceptance criterion (expected result)", "Status", "Tester", "Date"],
      tests, [9, 9, 80, 12, 16, 12])
ws_t = wb["UAT Tests"]
dv = DataValidation(type="list", formula1='"Not run,Pass,Fail,Blocked"', allow_blank=False)
ws_t.add_data_validation(dv)
last_t = len(tests) + 1
dv.add(f"D2:D{last_t}")
for row in ws_t.iter_rows(min_row=2, min_col=4, max_col=6):
    for c in row:
        c.font = INPUT
ws_t.conditional_formatting.add(f"D2:D{last_t}", CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor="E2F3EA")))
ws_t.conditional_formatting.add(f"D2:D{last_t}", CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", fgColor="FCE6E4")))
ws_t.conditional_formatting.add(f"D2:D{last_t}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", fgColor="FBF0D5")))

# Requirements traceability (test counts by formula across the story list)
req_rows = []
for r_ in C["requirements"]:
    req_rows.append([r_["id"], r_["type"], r_["text"], r_["priority"], r_["source"], r_["story"], None, None])
ws_r = sheet("Requirements", ["ID", "Type", "Requirement", "Priority (MoSCoW)", "Pain point", "User stories", "UAT tests", "Tests passed"],
             req_rows, [9, 14, 70, 12, 12, 16, 10, 12])
for i in range(2, len(req_rows) + 2):
    stories = req_rows[i - 2][5].split(", ")
    count = "+".join(f"COUNTIF('UAT Tests'!$B$2:$B${last_t},\"{s}\")" for s in stories)
    passed = "+".join(f"COUNTIFS('UAT Tests'!$B$2:$B${last_t},\"{s}\",'UAT Tests'!$D$2:$D${last_t},\"Pass\")" for s in stories)
    ws_r.cell(row=i, column=7, value=f"={count}").font = BODY
    ws_r.cell(row=i, column=8, value=f"={passed}").font = BODY
last_r = len(req_rows) + 1

# Summary
ws = wb.create_sheet("Summary", 1)
ws.column_dimensions["A"].width = 34
for col in "BCDE":
    ws.column_dimensions[col].width = 14
ws["A1"] = "Coverage summary"
ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
hdr = ["Priority", "Requirements", "With ≥1 test", "Fully passed", "% passed"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=3, column=j, value=h)
    c.fill, c.font, c.border = HEAD, HFONT, BOX
for k, pr in enumerate(["Must", "Should", "Could"], 4):
    ws.cell(row=k, column=1, value=pr)
    ws.cell(row=k, column=2, value=f'=COUNTIF(Requirements!$D$2:$D${last_r},A{k})')
    ws.cell(row=k, column=3, value=f'=COUNTIFS(Requirements!$D$2:$D${last_r},A{k},Requirements!$G$2:$G${last_r},">0")')
    ws.cell(row=k, column=4, value=f'=SUMPRODUCT((Requirements!$D$2:$D${last_r}=A{k})*(Requirements!$G$2:$G${last_r}>0)*(Requirements!$H$2:$H${last_r}=Requirements!$G$2:$G${last_r}))')
    ws.cell(row=k, column=5, value=f'=IFERROR(D{k}/B{k},0)')
ws.cell(row=7, column=1, value="Total")
for j, col in enumerate("BCD", 2):
    ws.cell(row=7, column=j, value=f"=SUM({col}4:{col}6)")
ws.cell(row=7, column=5, value="=IFERROR(D7/B7,0)")
for row in ws.iter_rows(min_row=4, max_row=7, max_col=5):
    for c in row:
        c.font, c.border = BODY, BOX
for c in ws[7]:
    c.font = Font(name=ARIAL, bold=True, size=10)
for k in range(4, 8):
    ws.cell(row=k, column=5).number_format = "0%"

ws["A10"] = "UAT status"
ws["A10"].font = Font(name=ARIAL, bold=True, size=11)
for k, st in enumerate(["Not run", "Pass", "Fail", "Blocked"], 11):
    ws.cell(row=k, column=1, value=st).font = BODY
    ws.cell(row=k, column=2, value=f"=COUNTIF('UAT Tests'!$D$2:$D${last_t},A{k})").font = BODY
ws.cell(row=15, column=1, value="Total test cases").font = Font(name=ARIAL, bold=True, size=10)
ws.cell(row=15, column=2, value="=SUM(B11:B14)").font = Font(name=ARIAL, bold=True, size=10)
ws["A17"] = "Go-live rule: every Must requirement fully passed (E4 = 100%)."
ws["A17"].font = Font(name=ARIAL, italic=True, size=10, color="5B6878")
ws["A18"] = '=IF(E4=1,"Ready for go-live decision","Not ready: Must requirements outstanding")'
ws["A18"].font = Font(name=ARIAL, bold=True, size=11)

wb.move_sheet("Summary", offset=-0)
out = ROOT / "docs" / "P2P_Requirements_Workbook.xlsx"
wb.save(out)
print("wrote", out.name, "and backlog/user_stories.csv")
